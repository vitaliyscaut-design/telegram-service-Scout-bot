import os
import logging
from pathlib import Path
from openpyxl import load_workbook
from dotenv import load_dotenv
load_dotenv()
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import (Application, CommandHandler, CallbackQueryHandler,
                          MessageHandler, ContextTypes, filters)

TOKEN = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
ADMIN_CHAT_ID = os.getenv("ADMIN_CHAT_ID", "").strip()
DATA_FILE = os.getenv("BOT_DATA_FILE", "Для Бота2-2.xlsx")
PAGE_SIZE = 8

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger(__name__)

def load_database():
    path = Path(DATA_FILE)
    if not path.exists():
        raise FileNotFoundError(f"Не знайдено файл бази: {path.resolve()}")
    ws = load_workbook(path, data_only=True, read_only=True).active
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        raise ValueError("Файл бази порожній")
    headers = [str(x or '').strip().lower() for x in all_rows[0]]
    aliases = {
        'model': ['модель трактора', 'модель'],
        'category': ['категорія', 'категория'],
        'problem': ['проблема'],
        'cause': ['причина/несправність', 'причина / несправність', 'несправність / можлива причина'],
        'detect': ['спосіб виявлення', 'способ выявления'],
        'fix': ['шлях усунення', 'способ устранения'],
        'self': ['самостійно', 'самостійний ремонт'],
        'instruction': [
            'інструкція для самостійного ремонту',
            'інструкція для самостійного ремонту:',
            'інструкція',
            'відео',
            'youtube',
            'посилання на відео'
        ],
        'master': ['майстер', 'виїзд майстра'],
        'service': ['сервіс', 'забір на сервіс'],
    }
    idx = {}
    for key, options in aliases.items():
        for i, h in enumerate(headers):
            if h in options:
                idx[key] = i
                break
        if key not in idx and key != 'instruction':
            raise ValueError(f"Не знайдено колонку для поля: {key}. Перевір заголовки таблиці.")

    data = []
    for line_no, row in enumerate(all_rows[1:], start=2):
        def val(key):
            if key not in idx:
                return ''
            i = idx[key]
            return str(row[i]).strip() if i < len(row) and row[i] is not None else ''
        model = val('model')
        category = val('category')
        problem = val('problem')
        cause = val('cause')
        if not (model and category and problem and cause):
            continue
        data.append({
            'id': len(data),
            'model': model,
            'category': category,
            'problem': problem,
            'cause': cause,
            'detect': val('detect'),
            'fix': val('fix'),
            'self': val('self'),
            'instruction': val('instruction'),
            'master': val('master'),
            'service': val('service'),
            'source_row': line_no,
        })
    if not data:
        raise ValueError("Не знайшов жодного повного рядка.")
    return data

DB = load_database()

def unique(values):
    return list(dict.fromkeys(values))

def page_markup(items, prefix, page=0):
    start = page * PAGE_SIZE
    subset = items[start:start + PAGE_SIZE]
    buttons = [[InlineKeyboardButton(label[:60], callback_data=f'{prefix}:{start + i}')]
               for i, label in enumerate(subset)]
    nav = []
    if page > 0:
        nav.append(InlineKeyboardButton('⬅️ Попередні', callback_data=f'page:{prefix}:{page - 1}'))
    if start + PAGE_SIZE < len(items):
        nav.append(InlineKeyboardButton('Наступні ➡️', callback_data=f'page:{prefix}:{page + 1}'))
    if nav:
        buttons.append(nav)
    buttons.append([InlineKeyboardButton('🏠 Головне меню', callback_data='menu')])
    return InlineKeyboardMarkup(buttons)

async def show_models(target, edit=False):
    models = unique(r['model'] for r in DB)
    markup = page_markup(models, 'model', 0)
    text = '🚜 *Діагностика трактора*\n\nОберіть модель трактора:'
    if edit:
        await target.edit_message_text(text, reply_markup=markup, parse_mode='Markdown')
    else:
        await target.reply_text(text, reply_markup=markup, parse_mode='Markdown')

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await show_models(update.message)

async def handle_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    data = q.data
    try:
        if data == 'menu':
            context.user_data.clear()
            await show_models(q, edit=True)
            return

        # ========== МОДЕЛЬ ==========
        if data.startswith('model:'):
            models = unique(r['model'] for r in DB)
            model = models[int(data.split(':')[1])]
            context.user_data['model'] = model
            categories = unique(r['category'] for r in DB if r['model'] == model)
            context.user_data['categories'] = categories
            await q.edit_message_text(
                f'🚜 *{model}*\n\nОберіть категорію:',
                reply_markup=page_markup(categories, 'cat', 0),
                parse_mode='Markdown'
            )
            return

        # ========== КАТЕГОРІЯ ==========
        if data.startswith('cat:'):
            categories = context.user_data.get('categories', [])
            category = categories[int(data.split(':')[1])]
            context.user_data['category'] = category
            model = context.user_data.get('model')
            problems = unique(r['problem'] for r in DB if r['model'] == model and r['category'] == category)
            context.user_data['problems'] = problems
            await q.edit_message_text(
                f'📂 *{category}*\n\nОберіть проблему:',
                reply_markup=page_markup(problems, 'prob', 0),
                parse_mode='Markdown'
            )
            return

        # ========== ПРОБЛЕМА ==========
        if data.startswith('prob:'):
            problems = context.user_data.get('problems', [])
            problem = problems[int(data.split(':')[1])]
            context.user_data['problem'] = problem
            model = context.user_data.get('model')
            category = context.user_data.get('category')
            causes = [r for r in DB if r['model'] == model and r['category'] == category and r['problem'] == problem]
            context.user_data['causes'] = causes
            labels = [r['cause'] for r in causes]
            await q.edit_message_text(
                f'⚠️ *{problem}*\n\nОберіть причину / несправність:',
                reply_markup=page_markup(labels, 'cause', 0),
                parse_mode='Markdown'
            )
            return

        # ========== ПРИЧИНА ==========
        if data.startswith('cause:'):
            causes = context.user_data.get('causes', [])
            record = causes[int(data.split(':')[1])]
            context.user_data['selected_record'] = record
            lines = [f"🔎 *Причина / несправність:*\n{record['cause']}"]
            if record['detect']:
                lines.append(f"\n🧪 *Спосіб виявлення:*\n{record['detect']}")
            if record['fix']:
                lines.append(f"\n🛠 *Шлях усунення:*\n{record['fix']}")
            buttons = []
            for key, label in [
                ('self', '🔧 Самостійний ремонт'),
                ('master', '👨‍🔧 Виклик майстра'),
                ('service', '🚜 Забір на сервіс')
            ]:
                if record[key].strip().lower() in ('так', 'так.', 'yes', 'да'):
                    buttons.append([InlineKeyboardButton(label, callback_data=f'route:{key}')])
            if not buttons:
                lines.append('\nℹ️ Для цього запису маршрут не позначений як доступний. Зверніться до сервісу 0800332034.')
            buttons.append([InlineKeyboardButton('⬅️ До причин', callback_data='back_causes')])
            buttons.append([InlineKeyboardButton('🏠 Головне меню', callback_data='menu')])
            await q.edit_message_text(
                '\n'.join(lines),
                reply_markup=InlineKeyboardMarkup(buttons),
                parse_mode='Markdown'
            )
            return

        # ========== НАЗАД ДО ПРИЧИН ==========
        if data == 'back_causes':
            problem = context.user_data.get('problem', '')
            causes = context.user_data.get('causes', [])
            await q.edit_message_text(
                f'⚠️ *{problem}*\n\nОберіть причину / несправність:',
                reply_markup=page_markup([r['cause'] for r in causes], 'cause', 0),
                parse_mode='Markdown'
            )
            return

        # ========== НАЗАД З МАРШРУТУ ==========
        if data == 'back_route':
            context.user_data['awaiting_contact'] = False
            record = context.user_data.get('selected_record')
            if not record:
                await show_models(q, edit=True)
                return
            lines = [f"🔎 *Причина / несправність:*\n{record['cause']}"]
            if record['detect']:
                lines.append(f"\n🧪 *Спосіб виявлення:*\n{record['detect']}")
            if record['fix']:
                lines.append(f"\n🛠 *Шлях усунення:*\n{record['fix']}")
            buttons = []
            for key, label in [
                ('self', '🔧 Самостійний ремонт'),
                ('master', '👨‍🔧 Виклик майстра'),
                ('service', '🚜 Забір на сервіс')
            ]:
                if record[key].strip().lower() in ('так', 'так.', 'yes', 'да'):
                    buttons.append([InlineKeyboardButton(label, callback_data=f'route:{key}')])
            if not buttons:
                lines.append('\nℹ️ Для цього запису маршрут не позначений як доступний. Зверніться до сервісу 0800332034.')
            buttons.append([InlineKeyboardButton('⬅️ До причин', callback_data='back_causes')])
            buttons.append([InlineKeyboardButton('🏠 Головне меню', callback_data='menu')])
            await q.edit_message_text(
                '\n'.join(lines),
                reply_markup=InlineKeyboardMarkup(buttons),
                parse_mode='Markdown'
            )
            return

        # ========== ПАГІНАЦІЯ ==========
        if data.startswith('page:'):
            _, prefix, page_s = data.split(':')
            page = int(page_s)
            if prefix == 'model':
                items = unique(r['model'] for r in DB)
                text = '🚜 *Діагностика трактора*\n\nОберіть модель трактора:'
            elif prefix == 'cat':
                items = context.user_data.get('categories', [])
                text = f"🚜 *{context.user_data.get('model', '')}*\n\nОберіть категорію:"
            elif prefix == 'prob':
                items = context.user_data.get('problems', [])
                text = f"📂 *{context.user_data.get('category', '')}*\n\nОберіть проблему:"
            else:
                items = [r['cause'] for r in context.user_data.get('causes', [])]
                text = f"⚠️ *{context.user_data.get('problem', '')}*\n\nОберіть причину / несправність:"
            await q.edit_message_text(
                text,
                reply_markup=page_markup(items, prefix, page),
                parse_mode='Markdown'
            )
            return

        # ========== МАРШРУТ ==========
        if data.startswith('route:'):
            route = data.split(':', 1)[1]
            labels = {
                'self': 'Самостійний ремонт',
                'master': 'Виклик майстра',
                'service': 'Забір на сервіс'
            }
            record = context.user_data.get('selected_record')
            if not record:
                await show_models(q, edit=True)
                return
            context.user_data['pending_route'] = labels[route]
            context.user_data['awaiting_contact'] = True

            text = (
                f"✅ Обраний маршрут: <b>{labels[route]}</b>\n"
                f"Модель: <b>{record.get('model', '')}</b>\n\n"
                "Якщо вам не вдалось самостійно вирішити проблему, напишіть нам у <b>Viber</b> на номер:\n"
                "<code>+380689797080</code>\n"
                "Або наберіть номер Гарячої лінії 0800332034\n\n"
                "При зверненні вкажіть наступну інформацію:\n"
                "• ваше ім’я;\n"
                "• номер телефону;\n"
                "• модель трактора;\n"
                "• кількість мотогодин;\n"
                "• населений пункт.\n"
            )

            # Відео-інструкція тільки для самостійного ремонту
            if route == 'self' and record.get('instruction'):
                url = record['instruction'].strip()
                if url.startswith('http'):
                    text += f'\n📹 <a href="{url}">Інструкція для самостійного ремонту (відео)</a>\n'

            text += "\nЯкщо поки не хочете надсилати заявку — натисніть «Головне меню»."

            await q.edit_message_text(
                text,
                reply_markup=InlineKeyboardMarkup([
                    [InlineKeyboardButton('⬅️ Назад', callback_data='back_route')],
                    [InlineKeyboardButton('🏠 Головне меню', callback_data='menu')]
                ]),
                parse_mode='HTML',
                disable_web_page_preview=False
            )
            return

    except (IndexError, ValueError, KeyError) as e:
        log.exception('Callback error: %s', e)
        await q.edit_message_text('Не вдалося обробити вибір. Натисніть /start і спробуйте ще раз.')

async def handle_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not context.user_data.get('awaiting_contact'):
        await update.message.reply_text('Оберіть пункт у меню кнопками нижче або введіть /start.')
        return
    record = context.user_data.get('selected_record', {})
    route = context.user_data.get('pending_route', 'Маршрут не вказаний')
    user = update.effective_user
    details = update.message.text.strip()
    if ADMIN_CHAT_ID:
        message = (
            '📩 *Нова сервісна заявка*\n'
            f"Маршрут: {route}\n"
            f"Модель: {record.get('model', '')}\n"
            f"Категорія: {record.get('category', '')}\n"
            f"Проблема: {record.get('problem', '')}\n"
            f"Причина: {record.get('cause', '')}\n"
            f"Спосіб виявлення: {record.get('detect') or 'не вказано'}\n"
            f"Шлях усунення: {record.get('fix') or 'не вказано'}\n"
            f"Користувач: @{user.username or 'без username'} (ID {user.id})\n"
            f"Контактні дані: {details}"
        )
        try:
            await context.bot.send_message(
                chat_id=int(ADMIN_CHAT_ID),
                text=message,
                parse_mode='Markdown'
            )
        except Exception:
            log.exception('Could not send request to admin')
            await update.message.reply_text(
                'Не вдалося передати заявку сервісу. Спробуйте пізніше або зв’яжіться із сервісом напряму.'
            )
            return
        await update.message.reply_text('✅ Заявку передано сервісу. З вами зв’яжуться для уточнення деталей.')
    else:
        await update.message.reply_text(
            'Дані збережені в цій розмові, але передача заявки сервісу ще не налаштована. '
            'Адміністратор має додати ADMIN_CHAT_ID у налаштування.'
        )
    context.user_data['awaiting_contact'] = False

async def help_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text('Введіть /start, щоб відкрити діагностичне меню.')

def main():
    print("=" * 50)
    print("DEBUG: TELEGRAM_BOT_TOKEN from os.getenv =", repr(os.getenv("TELEGRAM_BOT_TOKEN")))
    print("DEBUG: TOKEN variable =", repr(TOKEN))
    print("DEBUG: All env keys containing TOKEN:", [k for k in os.environ if "TOKEN" in k.upper()])
    print("=" * 50)

    if not TOKEN:
        raise SystemExit('Вкажіть TELEGRAM_BOT_TOKEN у файлі .env або змінних середовища.')

    app = Application.builder().token(TOKEN).build()
    app.add_handler(CommandHandler('start', start))
    app.add_handler(CommandHandler('help', help_cmd))
    app.add_handler(CallbackQueryHandler(handle_callback))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text))
    log.info('Loaded %s diagnostic rows', len(DB))
    app.run_polling(allowed_updates=Update.ALL_TYPES)

if __name__ == '__main__':
    main()
